"""PPTX 抽取器（python-pptx）。

按幻灯片编号保留定位；表格转成 `a | b | c` 行；备注页一并抽取，
因为很多课件的讲解内容只在 notes 里。
"""

from __future__ import annotations

from typing import Any

from ..winpath import isfile as _isfile, long_path as _long_path
from .base import EXTRACT_EMPTY, EXTRACT_FAILED, Segment, ExtractionResult, finalize


class PptxExtractor:
    name = "pptx"
    extensions = ("pptx", "potx")

    def can_handle(self, extension: str, media_type: str | None = None) -> bool:
        ext = (extension or "").lower().lstrip(".")
        if ext in self.extensions:
            return True
        return "presentationml" in (media_type or "").lower()

    def extract(self, path: str, *, max_chars: int | None = None,
                long_doc_threshold: int = 50000) -> ExtractionResult:
        if not _isfile(path):
            return finalize("", (), self.name, status=EXTRACT_FAILED,
                            error=f"文件不存在：{path}")
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            return finalize("", (), self.name, status=EXTRACT_FAILED,
                            error=f"python-pptx 不可用：{exc}")

        try:
            prs = Presentation(_long_path(path))
        except Exception as exc:  # noqa: BLE001
            return finalize("", (), self.name, status=EXTRACT_FAILED,
                            error=f"打开 PPTX 失败：{type(exc).__name__}: {exc}")

        parts: list[str] = []
        segments: list[Segment] = []
        slide_count = 0
        table_count = 0
        picture_count = 0
        try:
            for slide_index, slide in enumerate(prs.slides, start=1):
                slide_count += 1
                slide_lines: list[str] = []
                for shape in slide.shapes:
                    shape_type = str(getattr(shape, "shape_type", "")) or ""
                    if "PICTURE" in shape_type.upper():
                        picture_count += 1
                    if getattr(shape, "has_text_frame", False):
                        for para in shape.text_frame.paragraphs:
                            line = "".join(run.text or "" for run in para.runs).strip()
                            if not line:
                                line = (para.text or "").strip()
                            if line:
                                slide_lines.append(line)
                    if getattr(shape, "has_table", False):
                        table_count += 1
                        rows: list[list[str]] = []
                        for row in shape.table.rows:
                            rows.append([(c.text or "").strip() for c in row.cells])
                        flat = "\n".join(" | ".join(r) for r in rows)
                        if flat.strip():
                            slide_lines.append(flat)
                            segments.append(Segment(
                                text=flat, kind="table",
                                locator={"slide": slide_index, "rows": len(rows),
                                         "cols": max((len(r) for r in rows), default=0)},
                            ))
                notes_text = ""
                try:
                    if slide.has_notes_slide:
                        notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
                except Exception:  # noqa: BLE001
                    notes_text = ""

                body = "\n".join(slide_lines)
                if body.strip():
                    parts.append(f"[Slide {slide_index}]\n{body}")
                    segments.append(Segment(
                        text=body, kind="slide",
                        locator={"slide": slide_index, "char_count": len(body)},
                    ))
                if notes_text:
                    parts.append(f"[Slide {slide_index} notes]\n{notes_text}")
                    segments.append(Segment(
                        text=notes_text, kind="slide_notes",
                        locator={"slide": slide_index, "part": "notes"},
                    ))
        except Exception as exc:  # noqa: BLE001
            return finalize("\n".join(parts), tuple(segments), self.name,
                            status=EXTRACT_FAILED,
                            error=f"遍历 PPTX 失败：{type(exc).__name__}: {exc}",
                            extra_stats={"slide_count": slide_count})

        core: dict[str, Any] = {}
        try:
            cp = prs.core_properties
            core = {k: (str(getattr(cp, k) or "")[:200])
                    for k in ("title", "author", "subject", "keywords", "comments")}
        except Exception:  # noqa: BLE001
            core = {}

        stats = {
            "slide_count": slide_count,
            "table_count": table_count,
            "picture_count": picture_count,
            "core_properties": core,
        }
        text = "\n".join(parts)
        if not text.strip():
            return finalize("", tuple(segments), self.name, page_count=slide_count,
                            status=EXTRACT_EMPTY, error="PPTX 无可抽取文字",
                            extra_stats=stats)
        return finalize(text, segments, self.name, page_count=slide_count,
                        long_doc_threshold=long_doc_threshold, extra_stats=stats)
