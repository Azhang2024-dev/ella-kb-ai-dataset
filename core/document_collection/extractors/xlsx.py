"""表格类抽取器：XLSX / XLS / CSV。

- XLSX 用 openpyxl 的 read_only + data_only，只取计算结果不取公式；
- 每个 Sheet 一个 segment，locator 带 sheet 名与单元格范围（需求 3.1）；
- 行数设上限，超出就截断并在 stats 里明确记录，绝不静默丢数据；
- CSV 用标准库 csv，先嗅探分隔符，再复用文本抽取器的多编码解码。
"""

from __future__ import annotations

import csv
import io
import os
from typing import Any

from ..winpath import isfile as _isfile, long_path as _long_path, open_read as _open_read
from .base import EXTRACT_EMPTY, EXTRACT_FAILED, Segment, ExtractionResult, finalize
from .text import TextExtractor

MAX_ROWS_PER_SHEET = 20000
MAX_CELL_CHARS = 4000


class TabularExtractor:
    name = "tabular"
    extensions = ("xlsx", "xlsm", "xls", "csv", "tsv")

    def __init__(self) -> None:
        self._text = TextExtractor()

    def can_handle(self, extension: str, media_type: str | None = None) -> bool:
        ext = (extension or "").lower().lstrip(".")
        if ext in self.extensions:
            return True
        mt = (media_type or "").lower()
        return "spreadsheetml" in mt or mt in ("text/csv", "application/csv")

    # -- XLSX ------------------------------------------------------------
    def _extract_xlsx(self, path: str, long_doc_threshold: int) -> ExtractionResult:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            return finalize("", (), self.name, status=EXTRACT_FAILED,
                            error=f"openpyxl 不可用：{exc}")
        try:
            wb = load_workbook(_long_path(path), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            return finalize("", (), self.name, status=EXTRACT_FAILED,
                            error=f"打开工作簿失败：{type(exc).__name__}: {exc}")

        parts: list[str] = []
        segments: list[Segment] = []
        sheet_stats: list[dict[str, Any]] = []
        try:
            for sheet_name in wb.sheetnames:
                try:
                    ws = wb[sheet_name]
                except Exception as exc:  # noqa: BLE001
                    sheet_stats.append({"sheet": sheet_name, "error": str(exc)[:200]})
                    continue
                rows_text: list[str] = []
                row_count = 0
                truncated = False
                max_col = 0
                for row in ws.iter_rows(values_only=True):
                    row_count += 1
                    if row_count > MAX_ROWS_PER_SHEET:
                        truncated = True
                        break
                    cells = []
                    for value in row:
                        if value is None:
                            cells.append("")
                            continue
                        text = str(value)
                        if len(text) > MAX_CELL_CHARS:
                            text = text[:MAX_CELL_CHARS] + "…[TRUNCATED]"
                        cells.append(text.replace("\n", " ").strip())
                    max_col = max(max_col, len(cells))
                    if any(cells):
                        rows_text.append(" | ".join(cells))
                body = "\n".join(rows_text)
                if body.strip():
                    parts.append(f"[Sheet {sheet_name}]\n{body}")
                    segments.append(Segment(
                        text=body, kind="sheet",
                        locator={"sheet": sheet_name,
                                 "range": f"A1:{_col_letter(max_col)}{row_count}",
                                 "rows": len(rows_text), "cols": max_col,
                                 "truncated": truncated},
                    ))
                sheet_stats.append({"sheet": sheet_name, "rows_read": row_count,
                                    "rows_with_content": len(rows_text), "cols": max_col,
                                    "truncated": truncated})
        finally:
            try:
                wb.close()
            except Exception:  # noqa: BLE001
                pass

        stats = {"sheet_count": len(sheet_stats), "sheets": sheet_stats}
        text = "\n".join(parts)
        if not text.strip():
            return finalize("", tuple(segments), self.name, status=EXTRACT_EMPTY,
                            error="工作簿没有可抽取的单元格内容", extra_stats=stats)
        return finalize(text, segments, self.name, long_doc_threshold=long_doc_threshold,
                        extra_stats=stats)

    # -- CSV -------------------------------------------------------------
    def _extract_csv(self, path: str, long_doc_threshold: int) -> ExtractionResult:
        try:
            with _open_read(path, "rb") as fh:
                data = fh.read()
        except OSError as exc:
            return finalize("", (), self.name, status=EXTRACT_FAILED, error=f"读取失败：{exc}")

        text, encoding, decode_stats = self._text.decode_bytes(data)
        if not text.strip():
            return finalize("", (), self.name, status=EXTRACT_EMPTY,
                            error="CSV 为空", extra_stats={"encoding": encoding, **decode_stats})

        sample = text[:8192]
        delimiter = ","
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = "\t" if sample.count("\t") > sample.count(",") else ","

        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        rows_text: list[str] = []
        row_count = 0
        truncated = False
        max_col = 0
        try:
            for row in reader:
                row_count += 1
                if row_count > MAX_ROWS_PER_SHEET:
                    truncated = True
                    break
                cells = [(c or "").replace("\n", " ").strip()[:MAX_CELL_CHARS] for c in row]
                max_col = max(max_col, len(cells))
                if any(cells):
                    rows_text.append(" | ".join(cells))
        except csv.Error as exc:
            return finalize("\n".join(rows_text), (), self.name, status=EXTRACT_FAILED,
                            error=f"CSV 解析失败：{exc}",
                            extra_stats={"encoding": encoding, "delimiter": delimiter})

        body = "\n".join(rows_text)
        stats = {"encoding": encoding, "delimiter": delimiter, "rows_read": row_count,
                 "rows_with_content": len(rows_text), "cols": max_col,
                 "truncated": truncated, **decode_stats}
        segments = [Segment(text=body, kind="sheet",
                            locator={"sheet": os.path.basename(path),
                                     "range": f"row1-{row_count}", "rows": len(rows_text),
                                     "cols": max_col, "truncated": truncated,
                                     "delimiter": delimiter})] if body.strip() else []
        if not body.strip():
            return finalize("", (), self.name, status=EXTRACT_EMPTY,
                            error="CSV 没有有效数据行", extra_stats=stats)
        return finalize(body, segments, self.name, long_doc_threshold=long_doc_threshold,
                        extra_stats=stats)

    # -- 入口 ------------------------------------------------------------
    def extract(self, path: str, *, max_chars: int | None = None,
                long_doc_threshold: int = 50000) -> ExtractionResult:
        if not _isfile(path):
            return finalize("", (), self.name, status=EXTRACT_FAILED,
                            error=f"文件不存在：{path}")
        ext = os.path.splitext(path)[1].lower().lstrip(".")
        if ext in ("csv", "tsv"):
            return self._extract_csv(path, long_doc_threshold)
        if ext == "xls":
            return finalize("", (), self.name, status="conversion_required",
                            error="旧版 .xls 需要批准转换器，本期标记 conversion_required")
        return self._extract_xlsx(path, long_doc_threshold)


def _col_letter(index: int) -> str:
    """1 -> A, 26 -> Z, 27 -> AA。"""
    if index <= 0:
        return "A"
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters
