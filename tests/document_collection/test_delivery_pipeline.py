"""交付口径重构（2026-09-15）的专项测试锚点。

覆盖三处新行为，都是历史缺陷的直接防线：
1. **学科范围门禁**（screening）：模型判定 subject ∉ P0∪P1 → 淘汰；unknown → needs_review。
   历史：subject 只作展示，任何语种达标文档都能 accepted（180 份里 55% 学科越界）。
2. **query 句式多样化**（query_generation）：同一文档同开场套路重复 >3 次 → 报错 + 可重试修复。
   历史：实测"请根据文档内容"33 次，句式高度模板化。
3. **交付选取 + 真实名落盘**（delivery）：语种以正文判定、学科须属 P0∪P1、类型排排除项；
   真实文件名硬链接落地、交付表一行一条 query。
"""

from __future__ import annotations

import json
import os

from core.document_collection import contracts as C
from core.document_collection.query_generation import (
    QueryOptions,
    opening_signature,
    validate_queries,
)
from core.document_collection.screening import ScreeningOptions, semantic_review
from core.document_collection.delivery import (
    DeliverOptions,
    infer_material_type,
    select_documents,
)
from core.document_collection.state import StateStore


# ---------------------------------------------------------------------------
# 1. 学科范围门禁
# ---------------------------------------------------------------------------
def _accepted_decision():
    return {"status": "accepted", "reason_codes": [], "evidence": {}, "hard_failed": False}


class _ClassifyLLM:
    available = True
    name = "classify"

    def __init__(self, verdict):
        self._verdict = verdict

    def classify_document(self, document_ref, policy):
        return dict(self._verdict)


def _verdict(subject, *, education="university", queryability="high", evidence=None):
    return {
        "education_level": education, "subject": subject, "queryability": queryability,
        "exclude_kind": None, "evidence": evidence if evidence is not None else ["原文依据"],
        "confidence": 0.9, "needs_human_review": False,
    }


def _gate_opts():
    # 明确给出允许集合（等价 run_screen 从策略填充的效果）
    return ScreeningOptions(
        allowed_subjects=("law", "education_linguistics", "clinical_medicine_nursing",
                          "basic_science", "economics_management_accounting_finance",
                          "computer_science_it", "mathematics_statistics"),
        subject_gate=True, subject_unknown_action="needs_review",
    )


def test_out_of_scope_subject_is_rejected(real_policy):
    """心理学等范围外学科必须淘汰，而不是照单全收（历史缺陷 R4）。"""
    out = semantic_review(decision=_accepted_decision(),
                          llm=_ClassifyLLM(_verdict("psychology")),
                          document_ref={}, policy=real_policy, opts=_gate_opts())
    assert out["status"] == "rejected"
    assert C.RC_SUBJECT_OUT_OF_SCOPE in out["reason_codes"]


def test_in_scope_subject_is_accepted(real_policy):
    out = semantic_review(decision=_accepted_decision(),
                          llm=_ClassifyLLM(_verdict("law")),
                          document_ref={}, policy=real_policy, opts=_gate_opts())
    assert out["status"] == "accepted"
    assert C.RC_SUBJECT_OUT_OF_SCOPE not in out["reason_codes"]


def test_unknown_subject_goes_to_review_not_reject(real_policy):
    """判不出学科不等于越界，不能当越界误杀。"""
    out = semantic_review(decision=_accepted_decision(),
                          llm=_ClassifyLLM(_verdict("unknown")),
                          document_ref={}, policy=real_policy, opts=_gate_opts())
    assert out["status"] == "needs_review"
    assert C.RC_SUBJECT_OUT_OF_SCOPE in out["reason_codes"]


def test_subject_gate_can_be_disabled(real_policy):
    """显式关闭门禁时，范围外学科不再淘汰（旧批次对照用）。"""
    out = semantic_review(decision=_accepted_decision(),
                          llm=_ClassifyLLM(_verdict("psychology")),
                          document_ref={}, policy=real_policy,
                          opts=ScreeningOptions(subject_gate=False))
    assert out["status"] == "accepted"


def test_unknown_subject_action_reject(real_policy):
    out = semantic_review(decision=_accepted_decision(),
                          llm=_ClassifyLLM(_verdict("unknown")),
                          document_ref={}, policy=real_policy,
                          opts=ScreeningOptions(allowed_subjects=("law",),
                                                subject_gate=True,
                                                subject_unknown_action="reject"))
    assert out["status"] == "rejected"


def test_policy_exposes_subject_gate_defaults(real_policy):
    assert real_policy.subject_gate_enabled is True
    assert real_policy.subject_unknown_action in ("needs_review", "reject")


# ---------------------------------------------------------------------------
# 2. query 句式多样化
# ---------------------------------------------------------------------------
def _zh_queries(opening: str, n: int = 10) -> list[dict]:
    return [
        {"seq": i + 1, "type_key": "multiple_choice", "question_count": 5,
         "language": "zh", "text": f"{opening}第{i + 1}道题。",
         "evidence_locator": {"page": 1}, "compatibility_only": False}
        for i in range(n)
    ]


def test_opening_signature_groups_same_opener():
    assert opening_signature("请根据这份材料出5道题") == "请根据"
    assert opening_signature("给我出5道选择题") == "给我"
    assert opening_signature("帮我编3道填空题") == "帮我"
    assert opening_signature("根据这篇文章出题") == "根据"


def test_repeated_opener_is_an_error(real_policy):
    """10 条全用同一开场 → 报 query_style_repeated（历史实测 33 次雷同）。"""
    report = validate_queries(_zh_queries("请根据本文档，"),
                              policy=real_policy, document_language="zh",
                              opts=QueryOptions())
    assert report["ok"] is False
    assert any(e["code"] == C.RC_QUERY_STYLE_REPEATED for e in report["errors"])


def test_varied_openers_pass_style_check(real_policy):
    """开场各异（每签名 ≤3 次）应当通过句式检查。"""
    openers = ["给我出", "帮我编", "围绕本文", "用这份资料", "针对这篇",
               "根据这篇", "我需要", "麻烦按", "从这份材料", "这份文档能"]
    queries = []
    for i, op in enumerate(openers):
        queries.append({"seq": i + 1, "type_key": "multiple_choice" if i % 2 == 0 else "fill_in_blank",
                        "question_count": 5 if i % 2 == 0 else 3, "language": "zh",
                        "text": f"{op}出第{i + 1}道题。",
                        "evidence_locator": {"page": 1}, "compatibility_only": False})
    # 修正题型配比到策略要求，否则会被其它校验拦下——这里只关心句式维度
    report = validate_queries(queries, policy=real_policy, document_language="zh",
                              opts=QueryOptions())
    assert not any(e["code"] == C.RC_QUERY_STYLE_REPEATED for e in report["errors"])


def test_query_prompt_asks_for_varied_openings(real_policy):
    from core.document_collection.llm import build_query_prompt
    prompt = build_query_prompt({"language": "zh", "text": "正文"},
                                {**real_policy.query_plan, "max_tokens": 32000})
    assert "句式要求" in prompt
    assert "开头结构和句式必须各不相同" in prompt


# ---------------------------------------------------------------------------
# 3. 交付选取 + 真实名落盘
# ---------------------------------------------------------------------------
def _seed(store: StateStore, item_id: str, *, name: str, lang: str, subject: str,
          ext: str = "pdf", char_count: int = 5000, status: str = "accepted",
          exclude_kind=None, artifact_path=None) -> None:
    store.upsert_item({
        "item_id": item_id, "run_id": "r", "source_name": "src",
        "source_folder_path": "DOC-X", "item_token": item_id, "name": name,
        "item_type": "local_file", "extension": ext,
        "relative_path": f"DOC-X/{name}", "status": status, "stage": "screen",
        "metadata": {"relative_path": f"DOC-X/{name}"},
    })
    store.upsert_extraction(item_id, extractor="pdf", status="extracted",
                            char_count=char_count, language=lang, page_count=3)
    store.add_decision(item_id, "screen", "accepted",
                       ["subject=" + subject],
                       {"subject": subject, "education_level": "university",
                        "exclude_kind": exclude_kind})
    if artifact_path:
        store.upsert_artifact(item_id, path=artifact_path, sha256="a" * 64,
                              size_bytes=100, status="acquired")


def _store(workdir) -> StateStore:
    return StateStore(os.path.join(str(workdir), "s.sqlite3"))


def test_select_filters_out_of_scope_subject(real_policy, workdir):
    store = _store(workdir)
    try:
        _seed(store, "IN1", name="民法总论.pdf", lang="zh", subject="law")
        _seed(store, "OUT1", name="心理学导论.pdf", lang="zh", subject="psychology")
        store.commit()
        out = select_documents(policy=real_policy, store=store,
                               options=DeliverOptions(per_language=30))
        ids = [r["item_id"] for r in out["selected"]]
        assert ids == ["IN1"], ids
        assert any("subject_out_of_scope" in k for k in out["dropped"])
    finally:
        store.close()


def test_select_uses_detected_language_not_folder(real_policy, workdir):
    """语种以正文判定为准：zh 目录里的俄语文件不应进 zh 批次（历史 R5）。"""
    store = _store(workdir)
    try:
        _seed(store, "RU1", name="逻辑学的革命.txt", lang="ru", subject="mathematics_statistics")
        store.commit()
        out = select_documents(policy=real_policy, store=store,
                               options=DeliverOptions(per_language=30, languages=("zh",)))
        assert out["selected"] == []
        assert any("lang_not_target" in k for k in out["dropped"])
    finally:
        store.close()


def test_select_excludes_document_kinds(real_policy, workdir):
    store = _store(workdir)
    try:
        _seed(store, "R1", name="年度报告.pdf", lang="zh", subject="law",
              exclude_kind="annual_report")
        store.commit()
        out = select_documents(policy=real_policy, store=store,
                               options=DeliverOptions(per_language=30, languages=("zh",)))
        assert out["selected"] == []
        assert any("excluded_kind" in k for k in out["dropped"])
    finally:
        store.close()


def test_deliver_links_real_filenames_and_writes_xlsx(real_policy, workdir):
    from core.document_collection.delivery import run_deliver
    ws = str(workdir)
    src_dir = os.path.join(ws, "artifacts", "aa")
    os.makedirs(src_dir, exist_ok=True)
    src_file = os.path.join(src_dir, "a" * 64 + ".pdf")
    with open(src_file, "w", encoding="utf-8") as fh:
        fh.write("payload")

    store = _store(workdir)
    try:
        _seed(store, "Q1", name="民法总论.pdf", lang="zh", subject="law",
              artifact_path=src_file)
        store.upsert_query("Q1", 1, type_key="multiple_choice", question_count=5,
                           language="zh", payload={"text": "给我出5道选择题",
                                                   "type_key": "multiple_choice"},
                           status="validated", created_at="2026-09-15T00:00:00")
        store.commit()
        out = run_deliver(policy=real_policy, store=store, workspace_dir=ws,
                          run_id="r", options=DeliverOptions(per_language=30))
        assert out["selected_total"] == 1
        assert out["hardlinked"] + out["copied"] == 1
        # 真实文件名落地（不是哈希名）
        landed = os.path.join(ws, "delivery_files", "zh", "民法总论.pdf")
        assert os.path.isfile(landed), landed
        # 交付表存在且表头符合 sheet4
        xlsx = os.path.join(ws, "交付表_Query.xlsx")
        assert os.path.isfile(xlsx)
        from openpyxl import load_workbook
        wb = load_workbook(xlsx)
        headers = [c.value for c in wb.active[1]]
        assert headers == ["语种", "文件名", "格式", "字符数", "item_id", "类型",
                           "学科", "离线文件", "query", "query_type"]
        # 一行一条 query
        assert wb.active.max_row == 2
        assert wb.active.cell(2, 9).value == "给我出5道选择题"
    finally:
        store.close()


def test_infer_material_type():
    assert infer_material_type("2019期中试卷.pdf") == "试卷"
    assert infer_material_type("第3章习题.docx") == "作业习题"
    assert infer_material_type("物理实验指导.pdf") == "实验指导"
    assert infer_material_type("普通文档名.pdf") == ""


def test_material_type_head_classifier_and_folding():
    """资料类型：头部规则判真实体裁，再确定性折叠成五类（不放宽、不臆造）。"""
    from core.document_collection.material_type import (
        FOLD_TO_FIVE, FIVE_TYPES, fold_to_five, infer_from_head,
    )

    # 真实体裁
    assert infer_from_head("摘要：本文研究了…关键词：X")[0] == "学术论文"
    assert infer_from_head("ЭКЗАМЕНАЦИОННЫЙ БИЛЕТ № 1 по дисциплине")[0] == "试卷"
    assert infer_from_head("CHAPITRE 2 : BASES EN GÉNÉTIQUE")[0] == "教材"  # 大小写不敏感
    assert infer_from_head("实验一 一、实验目的 二、实验装置")[0] == "实验指导"
    # `cours` 不能命中英文 "discourse"（曾经的真实误判）
    assert infer_from_head("a discourse analysis of questions")[0] == ""
    # 文件名兜底：下划线分隔的分桶前缀
    assert infer_from_head("", "lecture_العقل_البشري_582_M.pdf")[0] == "讲义"

    # 折叠：严格五类，且只在五类里
    assert set(FOLD_TO_FIVE.values()) <= set(FIVE_TYPES) | {""}
    assert fold_to_five("学术论文") == "教材"
    assert fold_to_five("官方手册") == "教材"
    assert fold_to_five("课程大纲") == "讲义"
    assert fold_to_five("实验指导") == "作业习题"
    assert fold_to_five("其他") == ""          # 判不出留空，不硬套
    assert fold_to_five("不存在的体裁") == ""
    assert fold_to_five("课堂笔记") == "课堂笔记"   # 五类之一，必须可达


def test_material_type_detects_class_notes_without_false_positives():
    """课堂笔记必须判得出，且不能把参考文献/试卷说明里的 "lecture notes" 误吞。"""
    from core.document_collection.material_type import fold_to_five, infer_from_head

    # 标题处自述为笔记 → 命中（多语种）
    assert infer_from_head(
        "UMARU MUSA YAR'ADUA UNIVERSITY MTH 1311 LECTURE NOTE Prepared by…")[0] == "课堂笔记"
    assert infer_from_head(
        "ПОВОЛЖСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ КОНСПЕКТ ЛЕКЦИЙ ПО ДИСЦИПЛИНЕ")[0] == "课堂笔记"
    assert fold_to_five(infer_from_head("课堂笔记：第三章重点整理")[0]) == "课堂笔记"

    # 参考文献里的丛书名（正文尾部）不能命中
    ref = ("正文正文" * 2000) + " von Stechow (2008), Topics in Degree Semantics. "
    assert infer_from_head(ref + "Lecture Notes, Eberhard Karls Universität Tübingen.")[0] != "课堂笔记"
    # 试卷说明里的 "refer to your lecture notes" 是让考生看笔记，不是笔记正文
    exam = "## General Instructions\n• You can refer to your lecture notes.\n• Total 7 Marks"
    assert infer_from_head(exam)[0] != "课堂笔记"


def test_manifest_lists_deliver_capability():
    manifest = json.load(open(os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
        "skills", "document_collection", "capability-manifest.json"), encoding="utf-8"))
    names = [c["name"] for c in manifest["capabilities"]]
    assert "collection.deliver" in names
    assert "collection.deliver" in manifest["local_pipeline_order"]